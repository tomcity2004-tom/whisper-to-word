import streamlit as st
import whisper
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
import re
import io
import tempfile
import os

st.set_page_config(page_title="語音轉文字工具", layout="centered")
st.title("語音轉文字工具")
st.caption("上傳 MP3 / WAV / M4A，自動轉錄並排版成 Word 檔下載")

# ---------- 介面 ----------
uploaded_file = st.file_uploader(
    "上傳語音檔",
    type=["mp3", "wav", "m4a", "flac"],
    help="建議檔案大小不要太大，免費額度有限制"
)

col1, col2 = st.columns(2)
with col1:
    model_size = st.selectbox(
        "模型大小",
        ["tiny", "base", "small", "medium", "large"],
        index=2,
        help="越大越準但越慢、越耗資源。免費部署建議用 small 或 base"
    )
with col2:
    language = st.selectbox(
        "語言",
        ["zh", "en", None],
        format_func=lambda x: {"zh": "中文", "en": "英文", None: "自動偵測"}[x]
    )

pause_threshold = st.slider(
    "自動分段停頓閾值（秒）",
    min_value=0.5,
    max_value=3.0,
    value=1.2,
    step=0.1,
    help="停頓超過此秒數就換新段落"
)

# ---------- 快取模型 ----------
@st.cache_resource
def load_model(size):
    return whisper.load_model(size)

# ---------- 處理函式 ----------
def clean_text(text):
    return re.sub(r'\s+', ' ', text).strip()

def create_docx(paragraphs):
    doc = Document()

    # 頁面邊距
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 標題
    title = doc.add_heading("語音轉文字結果", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 段落
    for para_text in paragraphs:
        if not para_text:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.first_line_indent = Cm(0.74)

        run = p.add_run(para_text)
        run.font.name = "Microsoft JhengHei"
        run.font.size = Pt(12)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# ---------- 主流程 ----------
if uploaded_file is not None:
    if st.button("開始轉換", type="primary"):
        with st.spinner("正在載入模型並轉錄中，請稍候..."):
            # 暫存上傳檔案
            with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(uploaded_file.name)[1]) as tmp:
                tmp.write(uploaded_file.getvalue())
                audio_path = tmp.name

            try:
                model = load_model(model_size)

                result = model.transcribe(
                    audio_path,
                    language=language,
                    verbose=False,
                    initial_prompt="請使用中文標點符號，並保持自然的句子結構。"
                )

                # 依停頓自動分段
                paragraphs = []
                current_para = []
                prev_end = 0.0

                for seg in result["segments"]:
                    text = clean_text(seg["text"])
                    if not text:
                        continue
                    if current_para and (seg["start"] - prev_end > pause_threshold):
                        paragraphs.append("".join(current_para))
                        current_para = []
                    current_para.append(text)
                    prev_end = seg["end"]

                if current_para:
                    paragraphs.append("".join(current_para))

                if not paragraphs:
                    paragraphs = [clean_text(result["text"])]

                # 產生 Word 並提供下載
                docx_bytes = create_docx(paragraphs)

                st.success("轉換完成！")
                st.download_button(
                    label="下載 Word 檔",
                    data=docx_bytes,
                    file_name="語音轉文字結果.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            except Exception as e:
                st.error(f"發生錯誤：{e}")
            finally:
                if os.path.exists(audio_path):
                    os.unlink(audio_path)
else:
    st.info("請先上傳語音檔案")